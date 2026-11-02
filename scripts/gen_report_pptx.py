# -*- coding: utf-8 -*-
"""生成《人形机器人调试进展汇报》.pptx"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# 配色
ACCENT = RGBColor(0x1F, 0x4E, 0x79)   # 深蓝
ACCENT2 = RGBColor(0x2E, 0x75, 0xB6)  # 中蓝
LIGHT = RGBColor(0xF2, 0xF6, 0xFA)    # 浅底
GREY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
RED = RGBColor(0xC0, 0x39, 0x2B)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_rect(slide, x, y, w, h, color, line=False):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    if not line:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = color
    shp.shadow.inherit = False
    return shp


def add_text(slide, x, y, w, h, text, size=18, color=GREY, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = "微软雅黑"
    return tb


def title_bar(slide, title, idx):
    add_rect(slide, 0, 0, SW, Inches(1.15), ACCENT)
    add_rect(slide, 0, Inches(1.15), SW, Inches(0.06), ACCENT2)
    add_text(slide, Inches(0.6), Inches(0.18), Inches(11.5), Inches(0.8),
             title, size=28, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(12.2), Inches(0.3), Inches(0.9), Inches(0.6),
             f"{idx:02d}", size=22, color=RGBColor(0xBF, 0xD7, 0xEE),
             bold=True, align=PP_ALIGN.RIGHT)


def bullets(slide, items, x=Inches(0.8), y=Inches(1.6), w=Inches(11.7), size=18):
    tb = slide.shapes.add_textbox(x, y, w, SH - y - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        # it 支持 (level, text, color, bold) 或纯字符串
        if isinstance(it, tuple):
            lvl, txt = it[0], it[1]
            col = it[2] if len(it) > 2 and it[2] else GREY
            bd = it[3] if len(it) > 3 else False
        else:
            lvl, txt, col, bd = 0, it, GREY, False
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lvl
        p.space_after = Pt(10)
        prefix = "•  " if lvl == 0 else "–  "
        r = p.add_run()
        r.text = prefix + txt
        r.font.size = Pt(size if lvl == 0 else size - 2)
        r.font.bold = bd
        r.font.color.rgb = col if lvl == 0 else GREY
        r.font.name = "微软雅黑"
    return tb


# ---------- 第 1 页 封面 ----------
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, ACCENT)
add_rect(s, 0, Inches(4.55), SW, Inches(0.08), WHITE)
add_text(s, Inches(1), Inches(2.5), Inches(11.3), Inches(1.4),
         "人形机器人调试进展汇报", size=44, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(1), Inches(3.85), Inches(11.3), Inches(0.6),
         "行走稳定性优化阶段总结", size=22, color=RGBColor(0xD6, 0xE4, 0xF0),
         align=PP_ALIGN.CENTER)
add_text(s, Inches(1), Inches(4.9), Inches(11.3), Inches(0.6),
         "汇报人：________        日期：2026-10-28", size=16,
         color=RGBColor(0xBF, 0xD7, 0xEE), align=PP_ALIGN.CENTER)

# ---------- 第 2 页 概览 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "汇报概览", 2)
bullets(s, [
    (0, "目标：让机器人实现不搀扶下的快速、稳定行走", ACCENT, True),
    (0, "本轮已完成 3 项关键优化：", GREY, True),
    (1, "① 命令插值平滑 —— 抑制大阶跃、降低抖动"),
    (1, "② 零位 offset 校准 —— 修正结构偏差、改善右偏"),
    (1, "③ 更换行走模型 —— 适配全新电机"),
    (0, "当前状态：可独立快速稳定行走", GREEN, True),
    (0, "遗留问题：启停抖动、右转抖动（进行中）", RED, True),
])

# ---------- 第 3 页 工作一 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "工作一 · 命令插值平滑，降低大阶跃抖动", 3)
bullets(s, [
    (0, "现象：踝关节命令出现突发阶跃（峰值可达 ~127 rad/s），跟踪误差飙升、力矩塌陷，机体顿挫抖动", GREY),
    (0, "根因：RL 推理频率（50 Hz）与命令下发频率（200 Hz）存在间隙，原始动作在高频下发时形成突变台阶", GREY),
    (0, "措施：control 层对相邻帧命令做插值平滑，每 5 ms 一帧连续过渡，消除指令阶跃", ACCENT, True),
    (0, "效果：关节命令变化率受控，行走过程抖动明显下降", GREEN, True),
])

# ---------- 第 4 页 工作二 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "工作二 · 零位 offset 校准，改善右偏", 4)
bullets(s, [
    (0, "现象：直行/转向时机器人向右偏，存在结构性重心偏置", GREY),
    (0, "根因：装配与机械结构偏差导致左右腿零位不对称，右髋 roll 带载执行不足", GREY),
    (0, "措施：硬件配置中校准髋 roll 零位 offset，用不对称偏移补偿结构偏差", ACCENT, True),
], y=Inches(1.5))
# offset 数据小表
tbl_shape = s.shapes.add_table(3, 3, Inches(0.9), Inches(4.0), Inches(11.5), Inches(1.8))
tbl = tbl_shape.table
hdr = ["关节", "零位 offset", "作用"]
row = ["左髋 roll (left_hip_roll)", "+30.0°", "左腿零位基准"], \
      ["右髋 roll (right_hip_roll)", "−31.5°", "与左侧不对称，抵消右偏"]
for i, h in enumerate(hdr):
    c = tbl.cell(0, i); c.text = h
for ri, rv in enumerate([row[0], row[1]]):
    for ci, cv in enumerate(rv):
        c = tbl.cell(ri + 1, ci); c.text = cv
for ri in range(3):
    for ci in range(3):
        cell = tbl.cell(ri, ci)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if ci == 0 else PP_ALIGN.CENTER
        run = p.runs[0]
        run.font.size = Pt(15); run.font.name = "微软雅黑"
        cell.fill.solid()
        if ri == 0:
            run.font.bold = True; run.font.color.rgb = WHITE
            cell.fill.fore_color.rgb = ACCENT
        else:
            run.font.color.rgb = GREY
            cell.fill.fore_color.rgb = LIGHT if ri % 2 else WHITE

# ---------- 第 5 页 工作三 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "工作三 · 更换行走模型，适配全新电机", 5)
bullets(s, [
    (0, "背景：整机更换新一代电机（力矩系数 kt、控制模式等参数发生变化）", GREY),
    (0, "措施：更换/重训适配新电机的行走模型，同步更新关节 kp/kd、动作缩放与推理频率", ACCENT, True),
    (0, "效果：新模型与新电机输出特性匹配，控制精度与响应提升，支撑稳定快速行走", GREEN, True),
])

# ---------- 第 6 页 抖动改善数据表 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "优化效果 · 抖动降低对比", 6)
add_text(s, Inches(0.8), Inches(1.35), Inches(11.7), Inches(0.5),
         "关键指标优化前后对比（数据源自调试点日志观测，为代表性数值）", size=14, color=GREY)
data = [
    ["对比指标", "优化前", "优化后", "改善情况"],
    ["踝关节命令峰值变化率", "≈127 rad/s", "< 20 rad/s", "大幅下降"],
    ["行走机身抖动 / 顿挫", "频繁", "基本消除", "显著改善"],
    ["直行右偏", "明显右偏", "接近直线", "显著改善"],
    ["行走支撑方式", "需搀扶", "可独立快速行走", "达成目标"],
]
tw, th = Inches(11.7), Inches(3.4)
tbl_shape = s.shapes.add_table(len(data), 4, Inches(0.8), Inches(2.0), tw, th)
tbl = tbl_shape.table
tbl.columns[0].width = Inches(4.0)
tbl.columns[1].width = Inches(2.6)
tbl.columns[2].width = Inches(2.6)
tbl.columns[3].width = Inches(2.5)
for ri, rowdata in enumerate(data):
    for ci, val in enumerate(rowdata):
        cell = tbl.cell(ri, ci)
        cell.text = val
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if ci > 0 else PP_ALIGN.LEFT
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        run = p.runs[0]
        run.font.size = Pt(16 if ri == 0 else 15)
        run.font.name = "微软雅黑"
        if ri == 0:
            run.font.bold = True
            run.font.color.rgb = WHITE
            cell.fill.solid(); cell.fill.fore_color.rgb = ACCENT
        else:
            cell.fill.solid()
            cell.fill.fore_color.rgb = LIGHT if ri % 2 else WHITE
            run.font.color.rgb = GREEN if ci == 3 else GREY
            if ci == 3:
                run.font.bold = True

# ---------- 第 7 页 遗留问题 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "遗留问题", 7)
bullets(s, [
    (0, "⚠ 启停抖动", RED, True),
    (1, "现象：起步/停止瞬间明显抖动，甚至重心失撑"),
    (1, "初步分析：状态切换（如停止时 DAMP 零刚度瞬切）缺乏过渡，刚度突变"),
    (0, "⚠ 右转抖动", RED, True),
    (1, "现象：右转过程机体抖动、转向不顺畅"),
    (1, "初步分析：右髋 roll 带载执行不足，侧向重心转移受限（输出侧传动/执行能力）"),
])

# ---------- 第 8 页 后续计划 ----------
s = prs.slides.add_slide(BLANK)
title_bar(s, "后续计划", 8)
bullets(s, [
    (0, "启停抖动：为起停状态切换增加刚度/阻尼渐变插值，避免零刚度瞬切", ACCENT, True),
    (0, "右转抖动：进一步排查右髋 roll 带载执行能力（传动/参数），必要时结合命令速率限幅与零位微调", ACCENT, True),
    (0, "持续通过日志对比验证各项改动效果，保持配置基线稳定", GREY),
    (0, "里程碑目标：全场景（启停、直行、转向）平稳行走", GREEN, True),
])

out = "/home/zl/work/zl_robot/人形机器人调试进展汇报.pptx"
prs.save(out)
print("SAVED", out, "slides=", len(prs.slides._sldIdLst))
